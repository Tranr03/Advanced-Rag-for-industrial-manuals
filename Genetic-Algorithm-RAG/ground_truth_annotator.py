"""
Ground Truth Annotator
=======================
Reads Ground_Truth_Dataset_V2.xlsx, shows retrieved chunks per question,
and writes selected relevant_chunk_ids back into the Excel file.

Usage in notebook:
    from ground_truth_annotator import GroundTruthAnnotator
    from rag_retriever import RAGRetriever

    retriever = RAGRetriever()
    annotator = GroundTruthAnnotator(
        excel_path="Ground_Truth_Dataset_V2.xlsx",
        collection=retriever.collection,
        embedding_model=retriever.embedding_model,
    )
    annotator.status()
    annotator.run()
"""
import json
from openpyxl import load_workbook
from config import EMBEDDING_MODEL

try:
    from sentence_transformers import SentenceTransformer
except ImportError:
    SentenceTransformer = None


class GroundTruthAnnotator:

    def __init__(self, excel_path: str, collection, embedding_model=None,
                 top_n: int = 25):
        """
        Args:
            excel_path:      path to Ground_Truth_Dataset_V2.xlsx
            collection:      ChromaDB collection object (e.g. retriever.collection)
            embedding_model: SentenceTransformer instance; if None, loads EMBEDDING_MODEL
            top_n:           chunks to retrieve per question (default 25)
        """
        self.excel_path = excel_path
        self.collection = collection
        self.top_n = top_n

        if embedding_model is not None:
            self.embedding_model = embedding_model
        else:
            if SentenceTransformer is None:
                raise ImportError("sentence-transformers is required")
            print(f"Loading embedding model '{EMBEDDING_MODEL}'...")
            self.embedding_model = SentenceTransformer(EMBEDDING_MODEL)

        self.wb = load_workbook(excel_path)
        self.ws = self.wb.active
        self._parse_headers()

    # ------------------------------------------------------------------ #
    #  Excel helpers                                                       #
    # ------------------------------------------------------------------ #

    def _parse_headers(self):
        """Map column names (lowercase) to column indices from the header row."""
        headers = {}
        for col_idx, cell in enumerate(self.ws[1], start=1):
            if cell.value:
                headers[cell.value.strip().lower()] = col_idx

        self.col = {
            "id":                  headers.get("id"),
            "question":            headers.get("question"),
            "source_manual":       headers.get("source_manual"),
            "question_type":       headers.get("question_type"),
            "relevant_chunk_ids":  headers.get("relevant_chunk_ids"),
            "ground_truth_answer": headers.get("ground_truth_answer"),
            "notes":               headers.get("notes"),
            "checked":             headers.get("checked"),
        }

        missing = [k for k, v in self.col.items() if v is None]
        if missing:
            print(f"WARNING: Could not find columns: {missing}")
            print(f"Available headers: {list(headers.keys())}")

    def _get_cell(self, row, col_name):
        col_idx = self.col.get(col_name)
        if col_idx is None:
            return None
        return self.ws.cell(row=row, column=col_idx).value

    def _set_cell(self, row, col_name, value):
        col_idx = self.col.get(col_name)
        if col_idx is None:
            return
        self.ws.cell(row=row, column=col_idx).value = value

    def _get_unannotated_rows(self):
        """Return row indices where relevant_chunk_ids is empty."""
        rows = []
        for row_idx in range(2, self.ws.max_row + 1):
            val = self._get_cell(row_idx, "relevant_chunk_ids")
            if val is None or str(val).strip() == "":
                rows.append(row_idx)
        return rows

    # ------------------------------------------------------------------ #
    #  ChromaDB retrieval                                                  #
    # ------------------------------------------------------------------ #

    def _retrieve_chunks(self, query: str, source_manual=None) -> list[dict]:
        """
        Query ChromaDB using query_embeddings (matching the pipeline's approach).
        Collection uses hnsw:space=cosine so similarity = 1 - distance.
        """
        query_embedding = self.embedding_model.encode([query]).tolist()
        safe_n = min(self.top_n, self.collection.count())

        query_params = {
            "query_embeddings": query_embedding,
            "n_results":        safe_n,
            "include":          ["documents", "metadatas", "distances"],
        }

        if source_manual:
            source_manual = source_manual.replace(".pdf", "")
            query_params["where"] = {"source": {"$eq": source_manual}}

        try:
            raw = self.collection.query(**query_params)
        except Exception:
            query_params.pop("where", None)
            raw = self.collection.query(**query_params)

        # Also fall back if the filter matched nothing (not just on exception)
        if query_params.get("where") and not raw["ids"][0]:
            print(f"  [warn] source filter returned 0 results — retrying without filter")
            query_params.pop("where", None)
            raw = self.collection.query(**query_params)

        results = []
        for i in range(len(raw["ids"][0])):
            distance = raw["distances"][0][i]
            meta = raw["metadatas"][0][i] if raw["metadatas"] else {}
            results.append({
                "rank":         i + 1,
                "chunk_id":     raw["ids"][0][i],
                "content":      raw["documents"][0][i],
                "cosine_score": round(1 - distance, 4),
                "metadata":     meta,
            })
        return results

    def _keyword_search(self, keyword: str) -> list[dict]:
        """Python-side keyword search across all chunks (no embedding needed)."""
        all_data = self.collection.get(include=["documents", "metadatas"])
        kw = keyword.lower()
        found = []
        for i, doc in enumerate(all_data["documents"]):
            if kw in doc.lower():
                found.append({
                    "chunk_id": all_data["ids"][i],
                    "content":  doc[:120] + "...",
                    "metadata": all_data["metadatas"][i] if all_data["metadatas"] else {},
                })
        found = found[:20]
        print(f"\n  Found {len(found)} chunks containing '{keyword}':")
        for f in found:
            src  = f["metadata"].get("source", "?")
            page = f["metadata"].get("page", "?")
            print(f"    [{f['chunk_id']}] {src} p.{page} | {f['content']}")
        return found

    # ------------------------------------------------------------------ #
    #  Display helpers                                                     #
    # ------------------------------------------------------------------ #

    def _print_question_header(self, row_idx: int):
        q_id    = self._get_cell(row_idx, "id")
        question = self._get_cell(row_idx, "question")
        q_type  = self._get_cell(row_idx, "question_type")
        source  = self._get_cell(row_idx, "source_manual")
        answer  = self._get_cell(row_idx, "ground_truth_answer")
        notes   = self._get_cell(row_idx, "notes")

        print("\n" + "=" * 90)
        print(f"  [{q_id}] {question}")
        print(f"  Type: {q_type} | Source: {source}")
        print(f"  Expected Answer: {answer}")
        if notes:
            print(f"  Notes: {notes}")
        print("=" * 90)

    def _print_chunks(self, results: list[dict], max_preview: int = 150):
        print(f"\n  Retrieved {len(results)} chunks:\n")
        for r in results:
            preview = r["content"][:max_preview]
            if len(r["content"]) > max_preview:
                preview += "..."
            source = r["metadata"].get("source", "?")
            page   = r["metadata"].get("page", "?")
            ctype  = r["metadata"].get("type", "text")
            print(f"  [{r['rank']:>3}] Score: {r['cosine_score']:.4f} | "
                  f"Source: {source} | Page: {page} | Type: {ctype}")
            print(f"        {preview}")
            print(f"        ID: {r['chunk_id']}")
            print("-" * 90)

    def _show_full_chunk(self, results: list[dict], rank: int):
        if 1 <= rank <= len(results):
            r = results[rank - 1]
            print(f"\n{'=' * 60}")
            print(f"CHUNK RANK {rank} — ID: {r['chunk_id']}")
            print(f"Score: {r['cosine_score']} | "
                  f"Metadata: {json.dumps(r['metadata'], indent=2)}")
            print(f"\nFULL CONTENT:\n{r['content']}")
            print("=" * 60)
        else:
            print(f"  Invalid rank {rank}. Valid range: 1-{len(results)}")

    def _generate_llm_prompt(self, question: str, results: list[dict]):
        chunks_text = ""
        for r in results[:30]:
            chunks_text += f"\n--- CHUNK {r['rank']} (ID: {r['chunk_id']}) ---\n"
            chunks_text += (f"Source: {r['metadata'].get('source', '?')} | "
                            f"Page: {r['metadata'].get('page', '?')}\n")
            chunks_text += f"{r['content']}\n"

        prompt = (
            f"I am building a ground-truth dataset for evaluating a retrieval system.\n"
            f"Given the following question and retrieved text chunks from industrial manuals,\n"
            f"classify each chunk's relevance to answering the question.\n\n"
            f"QUESTION: {question}\n\n"
            f"For each chunk, respond with EXACTLY this format:\n"
            f"CHUNK <number>: <RELEVANT | PARTIALLY_RELEVANT | NOT_RELEVANT> "
            f"— <one sentence why>\n\n"
            f"CHUNKS:\n{chunks_text}\n"
            f"Classify all chunks now:"
        )
        print("\n" + "=" * 60)
        print("LLM PROMPT (copy-paste to Claude/ChatGPT):")
        print("=" * 60)
        print(prompt)
        print("=" * 60)

    # ------------------------------------------------------------------ #
    #  Main annotation loop                                               #
    # ------------------------------------------------------------------ #

    def run(self, filter_by_source: bool = True):
        """
        Iterate over unannotated questions, show chunks, write selections back
        to Excel. Saves after every question — safe to quit and resume anytime.

        Args:
            filter_by_source: narrow ChromaDB results to the question's
                              source_manual. Use 'nofilter' command at runtime
                              to override for a specific question.
        """
        unannotated = self._get_unannotated_rows()
        total = self.ws.max_row - 1
        done  = total - len(unannotated)

        print(f"Ground Truth Annotator")
        print(f"File:               {self.excel_path}")
        print(f"Total questions:    {total}")
        print(f"Already annotated:  {done}")
        print(f"Remaining:          {len(unannotated)}")
        print(f"\nCommands:")
        print(f"  1,3,7       → select chunks by rank")
        print(f"  show 5      → full content of rank 5")
        print(f"  keyword X   → keyword search across all chunks")
        print(f"  llm         → generate LLM pre-classification prompt")
        print(f"  nofilter    → re-retrieve without source_manual filter")
        print(f"  skip        → skip this question")
        print(f"  quit        → save and exit")

        for row_idx in unannotated:
            question = self._get_cell(row_idx, "question")
            source   = self._get_cell(row_idx, "source_manual")

            self._print_question_header(row_idx)

            source_filter = source if filter_by_source else None
            results = self._retrieve_chunks(question, source_filter)
            self._print_chunks(results)

            extra_chunk_ids = []

            while True:
                user_input = input(
                    "\n  >> Select relevant ranks (e.g. 1,3,7) or command: "
                ).strip()

                if not user_input:
                    continue

                if user_input.lower() == "quit":
                    self.wb.save(self.excel_path)
                    print(f"\n  Saved to {self.excel_path}. Exiting.")
                    return

                if user_input.lower() == "skip":
                    print("  Skipped.")
                    break

                if user_input.lower().startswith("show "):
                    try:
                        rank = int(user_input.split()[1])
                        self._show_full_chunk(results, rank)
                    except (ValueError, IndexError):
                        print("  Usage: show <rank_number>")
                    continue

                if user_input.lower().startswith("keyword "):
                    kw = user_input[8:].strip()
                    self._keyword_search(kw)
                    add = input("  Add chunk IDs from above? "
                                "(comma-separated IDs or Enter to skip): ").strip()
                    if add:
                        extra_chunk_ids.extend(
                            [x.strip() for x in add.split(",") if x.strip()]
                        )
                        print(f"  Queued {len(add.split(','))} extra chunk(s).")
                    continue

                if user_input.lower() == "llm":
                    self._generate_llm_prompt(question, results)
                    continue

                if user_input.lower() == "nofilter":
                    results = self._retrieve_chunks(question, None)
                    self._print_chunks(results)
                    continue

                # Parse rank numbers
                try:
                    ranks = [int(x.strip()) for x in user_input.split(",")]
                    selected_ids = []
                    for rank in ranks:
                        if 1 <= rank <= len(results):
                            selected_ids.append(results[rank - 1]["chunk_id"])
                        else:
                            print(f"  Warning: rank {rank} out of range "
                                  f"(1-{len(results)}), skipped.")

                    selected_ids.extend(extra_chunk_ids)

                    # Deduplicate, preserve order
                    seen: set = set()
                    unique_ids = []
                    for cid in selected_ids:
                        if cid not in seen:
                            seen.add(cid)
                            unique_ids.append(cid)

                    ids_string = ", ".join(unique_ids)
                    self._set_cell(row_idx, "relevant_chunk_ids", ids_string)
                    self._set_cell(row_idx, "checked", True)
                    self.wb.save(self.excel_path)

                    q_id = self._get_cell(row_idx, "id")
                    print(f"\n  Saved [{q_id}]: {len(unique_ids)} chunk(s) → {ids_string}")
                    break

                except ValueError:
                    print("  Invalid input. Enter rank numbers (e.g. 1,3,7) "
                          "or a command.")

        print(f"\n  All done! File saved: {self.excel_path}")

    # ------------------------------------------------------------------ #
    #  Progress reporting                                                  #
    # ------------------------------------------------------------------ #

    def status(self):
        """Print annotation progress broken down by manual and question type."""
        total = 0
        done  = 0
        by_manual: dict = {}
        by_type:   dict = {}

        for row_idx in range(2, self.ws.max_row + 1):
            total += 1
            val    = self._get_cell(row_idx, "relevant_chunk_ids")
            source = self._get_cell(row_idx, "source_manual") or "unknown"
            q_type = self._get_cell(row_idx, "question_type") or "unknown"
            is_done = val is not None and str(val).strip() != ""

            if is_done:
                done += 1

            by_manual.setdefault(source, {"total": 0, "done": 0})
            by_manual[source]["total"] += 1
            if is_done:
                by_manual[source]["done"] += 1

            by_type.setdefault(q_type, {"total": 0, "done": 0})
            by_type[q_type]["total"] += 1
            if is_done:
                by_type[q_type]["done"] += 1

        pct = (100 * done // total) if total else 0
        print(f"\n  Progress: {done}/{total} ({pct}%)\n")

        print("  By Manual:")
        for manual, c in by_manual.items():
            bar = "█" * c["done"] + "░" * (c["total"] - c["done"])
            print(f"    {str(manual)[:50]:<50} {bar}  {c['done']}/{c['total']}")

        print("\n  By Question Type:")
        for q_type, c in by_type.items():
            bar = "█" * c["done"] + "░" * (c["total"] - c["done"])
            print(f"    {q_type:<25} {bar}  {c['done']}/{c['total']}")
